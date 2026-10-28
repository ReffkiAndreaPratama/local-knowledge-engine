"""DOCX and XLSX parsers."""
from app.parsers.base import ParsedDocument, ParsedSection, ParseError


def parse_docx(path: str) -> ParsedDocument:
    """Extract paragraphs from a DOCX file."""
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ParseError("python-docx is not installed") from exc

    try:
        document = DocxDocument(path)
    except Exception as exc:
        raise ParseError(f"Cannot open DOCX: {exc}") from exc

    sections = []
    buffer = []
    current_heading = None

    def flush():
        text = "\n".join(buffer).strip()
        if text:
            sections.append(ParsedSection(text=text, heading=current_heading))
        buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            flush()
            current_heading = text
            buffer.append(text)
        else:
            buffer.append(text)
    flush()

    # Include table content
    for table_index, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sections.append(
                ParsedSection(text="\n".join(rows), heading=f"Table {table_index + 1}")
            )

    if not sections:
        raise ParseError("DOCX contains no readable text")

    return ParsedDocument(sections=sections)


def parse_xlsx(path: str) -> ParsedDocument:
    """Extract rows from an XLSX workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ParseError("openpyxl is not installed") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"Cannot open XLSX: {exc}") from exc

    sections = []
    try:
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    lines.append(" | ".join(values))
            if lines:
                sections.append(
                    ParsedSection(text="\n".join(lines), heading=sheet.title)
                )
    finally:
        workbook.close()

    if not sections:
        raise ParseError("XLSX contains no readable data")

    return ParsedDocument(sections=sections)
